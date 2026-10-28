"""JSON-driven SAT authoring engine. The original PTTEP DOCX is a visual reference,
not a document that receives text replacements. DOCX and PDF share one revision.
"""
from pathlib import Path
import subprocess, shutil, os, uuid
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .pttep_layout import configure_page, add_cover

ROOT=Path(__file__).resolve().parent
OUT=ROOT/"output"
BLUE=RGBColor(28,71,153)
def shade(cell, color="DCEAF8"):
    tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement("w:shd");sh.set(qn("w:fill"),color);tcPr.append(sh)
def write_table(doc,heads,rows):
    table=doc.add_table(rows=1,cols=len(heads))
    table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads):
        table.rows[0].cells[i].text=str(h);shade(table.rows[0].cells[i])
    trPr=table.rows[0]._tr.get_or_add_trPr()
    rep=OxmlElement("w:tblHeader");rep.set(qn("w:val"),"true");trPr.append(rep)
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v if v is not None else "")
    for row in table.rows:
        props=row._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit");props.append(cant)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                for run in paragraph.runs:run.font.size=Pt(9)
    return table
def keep_with_next(p):
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
def ensure_heading(doc, title, level):
    p=doc.add_heading(title,level=level)
    keep_with_next(p)
    return p
def add_image(doc,block):
    asset=block.get("asset","")
    source=ROOT/"assets"/asset
    if asset=="network_zms_ms.png" and source.is_file():
        doc.add_picture(str(source),width=Cm(min(float(block.get("widthCm",16)),17)))
        doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    else:
        raise FileNotFoundError("Referenced image has not been migrated or cannot be found: "+str(asset))
    if block.get("caption"):doc.add_paragraph(str(block["caption"]),style="Caption")
def add_block(doc,block):
    kind=block.get("type")
    if kind=="paragraph":doc.add_paragraph(str(block.get("text","")))
    elif kind=="heading":ensure_heading(doc,str(block.get("text","")),min(3,max(1,int(block.get("level",2)))))
    elif kind=="table":
        heads=block.get("columns",[])
        if heads:write_table(doc,heads,[[(row[i] if i<len(row) else "") for i in range(len(heads))] for row in block.get("rows",[])])
    elif kind=="image":add_image(doc,block)
    elif kind=="pageBreak":doc.add_page_break()
def create_document(data):
    doc=Document()
    styles=doc.styles
    styles["Normal"].font.name="Arial";styles["Normal"].font.size=Pt(10)
    for name in ["Heading 1","Heading 2","Heading 3"]:
        styles[name].font.name="Arial";styles[name].font.color.rgb=BLUE
    configure_page(doc,data)
    add_cover(doc,data)
    for section in data.get("sections",[]):
        if section.get("enabled",True) is False:continue
        ensure_heading(doc,str(section.get("title","Untitled Section")),level=1)
        for block in section.get("blocks",[]):add_block(doc,block)
    ensure_heading(doc,"SAT Test Register",level=1)
    write_table(doc,["Test ID","Description","Actual result","Status"],[
      [x.get("id",""),x.get("name",""),x.get("actual",""),x.get("status","NOT TESTED")] for x in data.get("tests",[])])
    ensure_heading(doc,"Acceptance & Sign-off",level=1)
    write_table(doc,["PTTEP","ADD VALUE SYSTEM"],[["Name: __________________","Name: __________________"],["Signature: ______________","Signature: ______________"],["Date: ___________________","Date: ___________________"]])
    OUT.mkdir(exist_ok=True)
    target=OUT/"PJ2601-0510_SAT_WORKING.docx";doc.save(target)
    return target
def render(data):
    target=create_document(data)
    pdf=target.with_suffix(".pdf")
    exe=shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        for folder in [os.environ.get("PROGRAMFILES",r"C:\Program Files"),os.environ.get("PROGRAMFILES(X86)",r"C:\Program Files (x86)")]:
            candidate=Path(folder)/"LibreOffice"/"program"/"soffice.exe"
            if candidate.is_file():exe=str(candidate);break
    if not exe:raise RuntimeError("DOCX created; LibreOffice is required for PDF preview")
    profile=(OUT/("lo_profile_"+uuid.uuid4().hex)).resolve().as_uri()
    result=subprocess.run([exe,"-env:UserInstallation="+profile,"--headless","--convert-to","pdf","--outdir",str(OUT),str(target)],capture_output=True,text=True,timeout=120)
    if result.returncode!=0 or not pdf.is_file():raise RuntimeError("PDF conversion failed: "+result.stderr[-500:])
    return {"docx":str(target),"pdf":str(pdf),"mode":"json-driven-authoring"}
