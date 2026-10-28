"""Conservative PTTEP SAT layout repair, preserving all existing paragraphs/tables/images."""
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def repair(source,destination):
    doc=Document(source)
    for i,sec in enumerate(doc.sections):
        if i>0:sec.start_type=WD_SECTION_START.CONTINUOUS
        sec.header_distance=Cm(0.85)
        if i==0:
            sec.top_margin=Cm(0.9)
            sec.bottom_margin=Cm(0.65)
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip():
            for br in paragraph._p.xpath('.//w:br[@w:type="page"]'):
                br.getparent().remove(br)
        if paragraph.style.name.startswith("Heading") or paragraph.text.strip()=="4.4 IP TABLE":
            paragraph.paragraph_format.keep_with_next=True
            paragraph.paragraph_format.keep_together=True
    for index,table in enumerate(doc.tables):
        if index==0 or len(table.columns)<=7:
            for row in table.rows:
                trPr=row._tr.get_or_add_trPr()
                if trPr.find(qn("w:cantSplit")) is None:
                    trPr.append(OxmlElement("w:cantSplit"))
    doc.save(destination)
    return destination
