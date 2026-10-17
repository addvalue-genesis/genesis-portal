"""PTTEP-style page chrome and controlled cover structure for JSON-driven SAT reports."""
from pathlib import Path
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent

def page_field(paragraph, code):
    run=paragraph.add_run()
    begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    text=OxmlElement("w:instrText");text.set(qn("xml:space"),"preserve");text.text=" "+code+" "
    sep=OxmlElement("w:fldChar");sep.set(qn("w:fldCharType"),"separate")
    val=OxmlElement("w:t");val.text="1"
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for el in (begin,text,sep,val,end):run._r.append(el)

def no_split(table):
    for row in table.rows:
        trpr=row._tr.get_or_add_trPr()
        item=OxmlElement("w:cantSplit");trpr.append(item)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.space_before=Pt(0)
                for run in p.runs:run.font.size=Pt(8)

def set_cell(cell,value,bold=False,center=True):
    cell.text=""
    p=cell.paragraphs[0]
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(str(value));r.bold=bold;r.font.size=Pt(8)

def logo_or_label(cell,filename,label,width=Cm(2.1)):
    logo=ROOT/"assets"/filename
    if logo.exists():
        p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo),width=width)
    else:set_cell(cell,label,True)

def configure_page(doc,data):
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.top_margin=Cm(3.25);sec.bottom_margin=Cm(1.8)
    sec.header_distance=Cm(0.8);sec.footer_distance=Cm(1.0)
    sec.different_first_page_header_footer=True
    header=sec.header
    table=header.add_table(rows=3,cols=3,width=Cm(17))
    table.style="Table Grid"
    table.autofit=False
    for row in table.rows:
        for i,w in enumerate((Cm(3.1),Cm(10.6),Cm(3.3))):row.cells[i].width=w
    logo_or_label(table.cell(0,0),"pttep_logo.png","PTTEP")
    set_cell(table.cell(0,1),"SITE ACCEPTANCE TEST (SAT)\nTYPE EXTERNAL ANTENNA",True)
    logo_or_label(table.cell(0,2),"addvalue_logo.png","ADD VALUE SYSTEM")
    set_cell(table.cell(1,0),"PTT Exploration and Production")
    set_cell(table.cell(1,1),data.get("project_title",""))
    set_cell(table.cell(1,2),"Add Value System Co., Ltd.")
    set_cell(table.cell(2,0),data.get("document_no","DOC NO. PENDING"))
    set_cell(table.cell(2,1),"REV "+str(data.get("revision","INTERNAL"))+"  |  STATUS: INTERNAL WORKING")
    p=table.cell(2,2).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Page ");page_field(p,"PAGE");p.add_run(" of ");page_field(p,"NUMPAGES")
    no_split(table)
    header.paragraphs[0].text=""
    # Cover uses its own document-control block and no running header.
    sec.first_page_header.paragraphs[0].text=""
    sec.footer.paragraphs[0].text=""
    sec.first_page_footer.paragraphs[0].text=""
    return sec

def add_cover(doc,data):
    title=doc.add_paragraph()
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before=Cm(2.4)
    r=title.add_run("ZAWTIKA FIELD DEVELOPMENT\nPROJECT SITE \"ZMS/MS\"")
    r.bold=True;r.font.size=Pt(19)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Cm(1)
    p.add_run(data.get("project_title","")).font.size=Pt(12)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Cm(2)
    r=p.add_run("SITE ACCEPTANCE TEST (SAT)\nTYPE EXTERNAL ANTENNA");r.bold=True;r.font.size=Pt(17)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Cm(3)
    p.add_run("SERVICE ORDER NO. "+str(data.get("po_number",""))).bold=True
    t=doc.add_table(rows=4,cols=3);t.style="Table Grid"
    content=[
      ["REVISION","DATE","ISSUE STATUS"],
      [str(data.get("revision","INTERNAL")),"","INTERNAL WORKING DRAFT"],
      ["DOCUMENT NO.",data.get("document_no","PENDING"),"NOT APPROVED"],
      ["PTTEP","PTT EXPLORATION AND PRODUCTION ZAWTIKA","ADD VALUE SYSTEM"]
    ]
    for i,row in enumerate(content):
        for j,val in enumerate(row):set_cell(t.cell(i,j),val,i==0)
    no_split(t)
    doc.add_page_break()
